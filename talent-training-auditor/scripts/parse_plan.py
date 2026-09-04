#!/usr/bin/env python3
"""
人才培养方案结构化解析器
支持 Word(.docx)、PDF、纯文本 格式
输出标准化的章节树结构，供后续审核和对比使用
"""

import sys
import json
import re
import os
from pathlib import Path


def extract_text_from_docx(file_path: str) -> str:
    """从 .docx 文件提取纯文本"""
    try:
        import docx
        doc = docx.Document(file_path)
        paragraphs = []
        for para in doc.paragraphs:
            if para.text.strip():
                paragraphs.append(para.text.strip())
        # 也提取表格内容
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)
        return "\n".join(paragraphs)
    except ImportError:
        # fallback: 使用系统命令
        result = os.popen(f"python3 -c \"import docx; doc=docx.Document('{file_path}'); [print(p.text) for p in doc.paragraphs if p.text.strip()]\" 2>/dev/null").read()
        return result if result else f"[无法解析 .docx 文件，请确认已安装 python-docx: {file_path}]"


def extract_text_from_pdf(file_path: str) -> str:
    """从 PDF 文件提取文本（依赖 pdf skill 的能力）"""
    try:
        import pdfplumber
        text_parts = []
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
        return "\n".join(text_parts)
    except ImportError:
        try:
            import PyPDF2
            text_parts = []
            with open(file_path, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                for page in reader.pages:
                    text_parts.append(page.extract_text() or "")
            return "\n".join(text_parts)
        except ImportError:
            return f"[无法解析 PDF 文件，建议先通过 pdf skill 将其转为文本: {file_path}]"


def load_file(file_path: str) -> str:
    """根据文件类型加载内容"""
    path = Path(file_path)
    if not path.exists():
        return f"[文件不存在: {file_path}]"

    suffix = path.suffix.lower()
    if suffix == ".docx":
        return extract_text_from_docx(file_path)
    elif suffix == ".pdf":
        return extract_text_from_pdf(file_path)
    elif suffix in (".txt", ".md", ""):
        return path.read_text(encoding="utf-8", errors="ignore")
    else:
        # 尝试作为文本读取
        try:
            return path.read_text(encoding="utf-8", errors="ignore")
        except Exception:
            return f"[不支持的文件格式: {suffix}]"


# 常见章节标题模式（适配高校人才培养方案）
SECTION_PATTERNS = [
    r"^[一二三四五六七八九十]+[、．.]\s*(.+)",          # 一、专业简介
    r"^第[一二三四五六七八九十]+[章节部分]\s*(.+)",      # 第一章 培养目标
    r"^\d+[、．.]\s*(.+)",                              # 1、培养目标
    r"^\d+\.\d+\s+(.+)",                               # 1.1 课程设置
    r"^（[一二三四五六七八九十]）\s*(.+)",               # （一）培养目标
    r"^[（(]\d+[）)]\s*(.+)",                           # （1）专业名称
]

# 人才培养方案核心模块关键词
MODULE_KEYWORDS = {
    "basic_info": ["专业名称", "专业代码", "学制", "学位", "授予", "层次", "基本信息"],
    "training_goals": ["培养目标", "培养定位", "人才培养目标"],
    "graduation_requirements": ["毕业要求", "毕业生", "能力要求", "素质要求", "知识要求"],
    "curriculum": ["课程体系", "课程设置", "课程结构", "必修课", "选修课", "公共基础课", "专业核心课"],
    "credit_hours": ["学分", "学时", "总学分", "总学时", "学分分配", "学时分配"],
    "practice": ["实践教学", "实习", "实训", "毕业设计", "毕业论文", "实验", "综合实践"],
    "teaching_plan": ["教学进程", "课程安排", "学期安排", "教学计划"],
    "quality_assurance": ["质量保障", "质量监控", "评价机制", "考核方式"],
    "faculty": ["师资", "教师队伍", "师资队伍", "双师型"],
    "conditions": ["办学条件", "实验室", "实训基地", "教学资源"],
}


def detect_module(text: str) -> str:
    """识别段落所属模块"""
    for module_key, keywords in MODULE_KEYWORDS.items():
        for kw in keywords:
            if kw in text:
                return module_key
    return "other"


def parse_plan_text(text: str) -> dict:
    """
    将方案文本解析为结构化字典
    返回格式：
    {
        "raw_text": "...",
        "sections": [
            {"title": "一、专业基本信息", "level": 1, "content": "...", "module": "basic_info"},
            ...
        ],
        "key_data": {
            "major_name": "...",
            "major_code": "...",
            "degree_length": "...",
            "degree_type": "...",
            "total_credits": "...",
            "total_hours": "...",
            "required_ratio": "...",
            "elective_ratio": "...",
            "practice_ratio": "...",
        }
    }
    """
    lines = text.split("\n")
    sections = []
    current_section = None
    current_content = []

    for line in lines:
        line = line.strip()
        if not line:
            continue

        is_heading = False
        heading_title = None
        heading_level = 2

        for pattern in SECTION_PATTERNS:
            m = re.match(pattern, line)
            if m:
                is_heading = True
                heading_title = line
                # 判断层级
                if re.match(r"^[一二三四五六七八九十]+[、．.]", line) or re.match(r"^第[一二三四五六七八九十]+[章节]", line):
                    heading_level = 1
                elif re.match(r"^\d+[、．.]", line):
                    heading_level = 2
                elif re.match(r"^\d+\.\d+", line):
                    heading_level = 3
                break

        if is_heading:
            # 保存上一个章节
            if current_section is not None:
                current_section["content"] = "\n".join(current_content).strip()
                sections.append(current_section)
            current_section = {
                "title": heading_title,
                "level": heading_level,
                "content": "",
                "module": detect_module(heading_title),
            }
            current_content = []
        else:
            current_content.append(line)
            if current_section:
                # 更新模块识别（基于内容）
                if current_section["module"] == "other":
                    mod = detect_module(line)
                    if mod != "other":
                        current_section["module"] = mod

    # 收尾
    if current_section is not None:
        current_section["content"] = "\n".join(current_content).strip()
        sections.append(current_section)
    elif current_content:
        # 没有识别到任何标题，整体作为一个节
        sections.append({
            "title": "（未识别章节结构）",
            "level": 1,
            "content": "\n".join(current_content),
            "module": "other"
        })

    # 提取关键数据
    key_data = extract_key_data(text)

    return {
        "raw_text": text,
        "sections": sections,
        "key_data": key_data,
        "section_count": len(sections),
        "total_chars": len(text),
    }


def extract_key_data(text: str) -> dict:
    """从文本中正则提取关键数值数据"""
    data = {}

    patterns = {
        "major_name": [r"专业名称[：:]\s*([^\n，,。；;]{2,20})"],
        "major_code": [r"专业代码[：:]\s*(\d{6,12})"],
        "degree_length": [r"学制[：:]\s*(\d+)\s*年", r"修业年限[：:]\s*(\d+)\s*年"],
        "degree_type": [r"学位类型[：:]\s*([^\n，,。；;]{2,15})", r"授予[^\n]*学位[：:]\s*([^\n，,。；;]{2,15})"],
        "total_credits": [r"总学分[：:]\s*(\d+\.?\d*)", r"毕业总学分[：:]\s*(\d+\.?\d*)"],
        "total_hours": [r"总学时[：:]\s*(\d+)", r"总课时[：:]\s*(\d+)"],
        "required_ratio": [r"必修[课程]*[比例占比]*[：:]\s*(\d+\.?\d*%?)"],
        "elective_ratio": [r"选修[课程]*[比例占比]*[：:]\s*(\d+\.?\d*%?)"],
        "practice_ratio": [r"实践[教学学时学分][比例占比]*[：:]\s*(\d+\.?\d*%?)", r"实践[教学]?[比例]*[：:]\s*(\d+\.?\d*%?)"],
    }

    for key, pats in patterns.items():
        for pat in pats:
            m = re.search(pat, text)
            if m:
                data[key] = m.group(1).strip()
                break

    return data


def main():
    if len(sys.argv) < 2:
        print(json.dumps({"error": "用法: parse_plan.py <文件路径或 'stdin'>", "sections": [], "key_data": {}}, ensure_ascii=False, indent=2))
        sys.exit(1)

    file_arg = sys.argv[1]

    if file_arg == "stdin" or file_arg == "-":
        text = sys.stdin.read()
    else:
        text = load_file(file_arg)

    result = parse_plan_text(text)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
